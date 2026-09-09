"""Round-trip и tampering для human-interface слоя. Требует openpyxl:
    uv run --group human-interface python -m unittest discover -s tests
Без группы тесты пропускаются с явной причиной, core-набор остаётся stdlib-only."""

import hashlib
import json
import tempfile
import unittest
from pathlib import Path

try:
    import openpyxl  # noqa: F401
    HAS_OPENPYXL = True
except ImportError:  # pragma: no cover
    HAS_OPENPYXL = False

from metrics.naturalness_ab import build_packet as build_ab_packet
from metrics.naturalness_ab import validate_responses as validate_ab
from metrics.naturalness_audit import build_packet as build_audit_packet
from metrics.naturalness_audit import validate_responses as validate_audit
from metrics.xlsx_interface import KINDS, collect_workbook, render_workbook


def audit_packet():
    stimuli = [
        {"package_id": "annotation-pilot-v0", "item_id": "pc-01", "language": "ru",
         "messages": [{"author": "a", "text": "Раз."}, {"author": "b", "text": "Два."}]},
        {"package_id": "annotation-pilot-v0", "item_id": "pn-01", "language": "ru",
         "messages": [{"author": "a", "text": "Одно сообщение."}]},
        {"package_id": "annotation-ux-v6", "item_id": "dg-04", "language": "ru",
         "messages": [{"author": "a", "text": "Три."}, {"author": "b", "text": "Четыре."}]},
    ]
    return build_audit_packet(stimuli, "auditor-1", "seed-1")


def ab_packet():
    candidates = {"schema_version": "rf.naturalness-ab-candidates.v1", "ab_id": "t", "items": [{
        "item_id": "pc-05", "language": "ru",
        "original": {"messages": [{"author": "a", "text": "A"}, {"author": "b", "text": "B orig"}], "target_index": 1},
        "candidates": [{"candidate_id": "pc-05-c1", "messages": [{"author": "a", "text": "A"}, {"author": "b", "text": "B one"}]},
                       {"candidate_id": "pc-05-c2", "messages": [{"author": "a", "text": "A"}, {"author": "b", "text": "B two"}]}],
    }]}
    return build_ab_packet(candidates, "rater-1", "seed-1")


def sha(packet: dict) -> str:
    return hashlib.sha256(json.dumps(packet, ensure_ascii=False).encode("utf-8")).hexdigest()


def roundtrip(wb):
    """Сохранить и заново открыть: то, что реально сделает человек."""
    from openpyxl import load_workbook
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "returned.xlsx"
        wb.save(path)
        return load_workbook(path)


@unittest.skipUnless(HAS_OPENPYXL, "openpyxl only in the human-interface group")
class AuditWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.packet, self.mapping = audit_packet()
        self.sha = sha(self.packet)
        self.wb = render_workbook("audit", self.packet, self.sha, "auditor-1")
        self.ws = self.wb["Оценка"]

    def fill(self, natural, invisible=None, other=""):
        for r in range(2, self.ws.max_row + 1):
            rid = self.ws.cell(row=r, column=1).value
            single = self.mapping["map"][rid]["n_messages"] == 1
            self.ws.cell(row=r, column=4).value = natural
            if not single:
                self.ws.cell(row=r, column=5).value = invisible
            self.ws.cell(row=r, column=6).value = other

    def test_render_hides_ids_and_prefills_not_applicable(self):
        self.assertTrue(self.ws.column_dimensions["A"].hidden)
        self.assertEqual(self.wb["_meta"].sheet_state, "hidden")
        self.assertTrue(self.ws.protection.sheet)
        dump = json.dumps([[c.value for c in row] for ws in self.wb for row in ws.iter_rows()], ensure_ascii=False)
        for token in ("pc-01", "pn-01", "dg-04", "stratum", "critic"):
            self.assertNotIn(token, dump)
        for r in range(2, self.ws.max_row + 1):
            rid = self.ws.cell(row=r, column=1).value
            cell = self.ws.cell(row=r, column=5)
            if self.mapping["map"][rid]["n_messages"] == 1:
                self.assertEqual(cell.value, "неприменимо")
                self.assertTrue(cell.protection.locked)
            else:
                self.assertFalse(cell.protection.locked)
            self.assertFalse(self.ws.cell(row=r, column=4).protection.locked)

    def test_roundtrip_to_canonical_jsonl(self):
        self.fill("отчасти", "нет", "слишком гладко")
        responses, issues = collect_workbook("audit", roundtrip(self.wb), self.packet, self.sha, "auditor-1")
        self.assertEqual(issues, [])
        self.assertEqual(validate_audit(responses, self.mapping["map"], "auditor-1"), [])
        by_id = {r["audit_item_id"]: r for r in responses}
        for rid, entry in self.mapping["map"].items():
            self.assertEqual(by_id[rid]["natural"], "partly")
            self.assertEqual(by_id[rid]["invisible_context"], "not_applicable" if entry["n_messages"] == 1 else "no")
            self.assertEqual(by_id[rid]["other_problem"], "слишком гладко")
            self.assertNotIn("note", by_id[rid])

    def test_labels_are_case_and_space_tolerant(self):
        self.fill(" Да ", "НЕТ")
        responses, issues = collect_workbook("audit", roundtrip(self.wb), self.packet, self.sha, "auditor-1")
        self.assertEqual(issues, [])
        self.assertTrue(all(r["natural"] == "yes" for r in responses))

    def test_wrong_packet_sha_is_refused(self):
        self.fill("да", "нет")
        _, issues = collect_workbook("audit", roundtrip(self.wb), self.packet, "0" * 64, "auditor-1")
        self.assertTrue(any("packet_sha256" in i for i in issues))

    def test_wrong_person_is_refused(self):
        self.fill("да", "нет")
        _, issues = collect_workbook("audit", roundtrip(self.wb), self.packet, self.sha, "auditor-2")
        self.assertTrue(any("один файл = один человек" in i for i in issues))

    def test_deleted_and_duplicated_rows(self):
        self.fill("да", "нет")
        last = self.ws.max_row
        for c in range(1, 8):  # дубль последней строки
            self.ws.cell(row=last + 1, column=c).value = self.ws.cell(row=last, column=c).value
        self.ws.delete_rows(2)  # удалить первую
        _, issues = collect_workbook("audit", roundtrip(self.wb), self.packet, self.sha, "auditor-1")
        joined = "\n".join(issues)
        self.assertIn("продублирована", joined)
        self.assertIn("удалена", joined)

    def test_value_outside_list_and_edited_text(self):
        self.fill("да", "нет")
        self.ws.cell(row=2, column=4).value = "наверное"
        self.ws.cell(row=3, column=3).value = "A: переписано"
        _, issues = collect_workbook("audit", roundtrip(self.wb), self.packet, self.sha, "auditor-1")
        joined = "\n".join(issues)
        self.assertIn("не из списка", joined)
        self.assertIn("текст переписки изменён", joined)

    def test_not_applicable_on_multi_message_is_caught_downstream(self):
        self.fill("да", "неприменимо")
        responses, issues = collect_workbook("audit", roundtrip(self.wb), self.packet, self.sha, "auditor-1")
        self.assertEqual(issues, [])
        self.assertTrue(any("multi-message item" in i for i in validate_audit(responses, self.mapping["map"], "auditor-1")))

    def test_renamed_header_is_structural_refusal(self):
        self.fill("да", "нет")
        self.ws.cell(row=1, column=4).value = "Норм?"
        _, issues = collect_workbook("audit", roundtrip(self.wb), self.packet, self.sha, "auditor-1")
        self.assertTrue(any("заголовки изменены" in i for i in issues))


@unittest.skipUnless(HAS_OPENPYXL, "openpyxl only in the human-interface group")
class AbWorkbookTests(unittest.TestCase):
    def test_roundtrip_and_no_original_side_in_workbook(self):
        packet, mapping = ab_packet()
        wb = render_workbook("ab", packet, sha(packet), "rater-1")
        ws = wb["Оценка"]
        dump = json.dumps([[c.value for c in row] for w in wb for row in w.iter_rows()], ensure_ascii=False)
        for token in ("pc-05", "original", "candidate", "left", "right"):
            self.assertNotIn(token, dump)
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=5).value = "вариант 2"
            ws.cell(row=r, column=6).value = "чуть отличается"
            ws.cell(row=r, column=7).value = "второй короче"
        responses, issues = collect_workbook("ab", roundtrip(wb), packet, sha(packet), "rater-1")
        self.assertEqual(issues, [])
        self.assertEqual(validate_ab(responses, mapping["map"], "rater-1"), [])
        self.assertTrue(all(r["more_natural"] == "right" and r["meaning_shift"] == "slight" and r["note"] == "второй короче"
                            for r in responses))
        self.assertEqual({r["pair_id"] for r in responses}, set(mapping["map"]))

    def test_empty_answer_is_refused(self):
        packet, _ = ab_packet()
        wb = render_workbook("ab", packet, sha(packet), "rater-1")
        _, issues = collect_workbook("ab", roundtrip(wb), packet, sha(packet), "rater-1")
        self.assertTrue(any("не из списка" in i for i in issues))


class KindSpecTests(unittest.TestCase):
    def test_labels_are_unique_and_comma_free(self):
        for spec in KINDS.values():
            for _, labels in spec["choice_columns"].values():
                self.assertEqual(len(set(labels.values())), len(labels))
                self.assertTrue(all("," not in v for v in labels.values()))


if __name__ == "__main__":
    unittest.main()
