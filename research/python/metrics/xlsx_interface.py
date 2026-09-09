"""Human-interface слой: дверь для испытуемого.

    canonical packet JSON  →  render  →  per-person .xlsx  →  человек заполняет
    →  returned .xlsx  →  collect  →  canonical response JSONL

Единственное место, где разрешён openpyxl (dependency group `human-interface`); core research
tooling остаётся stdlib-only. XLSX — derived artifact, source of truth — JSON packet: в книге
спрятаны packet_sha256, схема и opaque row ids, и collect отказывает, если книга сделана не из
того пакета, который лежит на диске.

В книге нет canonical item ids, strata, labels, флагов critic-1 и стороны оригинала — их нет
и в пакете. Человек редактирует только жёлтые ячейки ответов (лист защищён без пароля,
ответные ячейки разблокированы); выпадающие списки с закрытым набором значений; у одиночных
сообщений «невидимый контекст» уже стоит «неприменимо» и заблокирован.

collect валидирует: packet_sha256 и person совпадают; заголовки и листы не менялись; набор
opaque ids ровно тот, что в пакете (нет пропусков и дублей); только разрешённые значения;
not_applicable только там, где разрешён (через validate_responses соответствующего инструмента).
Только после этого пишется JSONL; уже существующий JSONL не перезаписывается никогда.
Returned .xlsx копируется как получен в responses/returned/ и получает sha256.

    uv run --group human-interface python -m metrics.xlsx_interface render  --kind audit --dir ../../data/pilot/naturalness-audit/v0-all
    uv run --group human-interface python -m metrics.xlsx_interface render  --kind ab    --dir ../../data/pilot/naturalness-ab/v0-flagged
    uv run --group human-interface python -m metrics.xlsx_interface collect --kind audit --dir ... --person auditor-1 --returned ~/Downloads/auditor-1.xlsx
"""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import sys
from pathlib import Path

WORKBOOK_SCHEMA = "rf.xlsx-interface.v1"
FONT_NAME = "Arial"
FILL_ANSWER = "FFF2CC"   # жёлтый: сюда пишет человек
FILL_HEADER = "D9E1F2"

# --- kinds -----------------------------------------------------------------------------------

KINDS: dict[str, dict] = {
    "audit": {
        "manifest": "audit-manifest.json",
        "persons_key": "auditors",
        "packet_key": "items",
        "row_id_key": "audit_item_id",
        "response_schema": "rf.naturalness-audit-response.v1",
        "response_id_key": "audit_item_id",
        "person_key": "auditor_id",
        "headers": ("id", "№", "Переписка", "Естественно?", "Невидимый контекст?", "Другая проблема?", "Комментарий"),
        "choice_columns": {  # column index (1-based) → (field, code→label)
            4: ("natural", {"yes": "да", "partly": "отчасти", "no": "нет"}),
            5: ("invisible_context", {"yes": "да", "no": "нет", "not_applicable": "неприменимо"}),
        },
        "text_columns": {6: "other_problem", 7: "note"},
        "instructions": [
            "Перед вами 46 коротких фрагментов переписки между двумя близкими людьми, A и B. Порядок случайный.",
            "По каждому фрагменту ответьте на три вопроса в жёлтых ячейках той же строки. В первых двух — выберите значение из списка; в последних двух можно писать словами.",
            "1. Естественно? — Звучит ли это как естественная переписка двух близких людей? да / отчасти / нет.",
            "2. Невидимый контекст? — Требует ли ответ B существенного невидимого контекста между A и B, сцены, которой нет в переписке? да / нет. Если во фрагменте одно сообщение, там уже стоит «неприменимо», менять не нужно.",
            "3. Другая проблема? — Есть ли другая наблюдаемая проблема формулировки? Если нет, оставьте пустым; если да, коротко словами.",
            "Комментарий — по желанию.",
            "Не ищите категории, не гадайте, «что здесь проверяют», не додумывайте историю пары. Отвечайте как читатель реальной переписки.",
            "Не переименовывайте листы и колонки, не удаляйте и не добавляйте строки, не меняйте текст переписок. Заполняйте только жёлтые ячейки.",
            "Закончили — сохраните файл и пришлите его обратно тому, кто его выдал. После отправки ответы не меняются. До отправки не обсуждайте фрагменты ни с кем.",
        ],
        "example": [("A: Выезжаю через десять минут.\nB: Купи хлеба по дороге, пожалуйста.", "да", "нет", "", ""),
                    ("A: Дождь обещают весь день, возьми зонт.", "да", "неприменимо", "", "")],
    },
    "ab": {
        "manifest": "ab-manifest.json",
        "persons_key": "raters",
        "packet_key": "pairs",
        "row_id_key": "pair_id",
        "response_schema": "rf.naturalness-ab-response.v1",
        "response_id_key": "pair_id",
        "person_key": "rater_id",
        "headers": ("id", "№", "Вариант 1", "Вариант 2", "Что естественнее?", "Изменился смысл или сила?", "Комментарий"),
        "choice_columns": {
            5: ("more_natural", {"left": "вариант 1", "right": "вариант 2", "no_difference": "без разницы"}),
            6: ("meaning_shift", {"same": "тот же", "slight": "чуть отличается", "substantial": "существенно отличается"}),
        },
        "text_columns": {7: "note"},
        "instructions": [
            "Перед вами 26 пар коротких переписок между A и B. В каждой паре два варианта одного и того же обмена; они отличаются одной-двумя репликами. Порядок случайный.",
            "По каждой паре ответьте на два вопроса в жёлтых ячейках той же строки, выбрав значение из списка.",
            "1. Что естественнее? — Какой из двух вариантов больше похож на реальную переписку пары? вариант 1 / вариант 2 / без разницы.",
            "2. Изменился смысл или сила? — Отличаются ли варианты по смыслу или по накалу сказанного? тот же / чуть отличается / существенно отличается. Это не про «какой лучше», а про «то же ли самое сказано и с той же ли силой». Вариант может быть естественнее и при этом говорить другое — это важно отметить.",
            "Комментарий — по желанию.",
            "Не ищите категории и не гадайте, какой вариант «исходный»: этого в задаче нет. Отвечайте по ощущению читателя реальной переписки.",
            "Не переименовывайте листы и колонки, не удаляйте и не добавляйте строки, не меняйте текст переписок. Заполняйте только жёлтые ячейки.",
            "Закончили — сохраните файл и пришлите его обратно тому, кто его выдал. После отправки ответы не меняются. До отправки не обсуждайте пары ни с кем.",
        ],
        "example": [("A: Во сколько ты сегодня дома?\nB: К восьми, наверное.", "A: Во сколько ты сегодня дома?\nB: Я планирую вернуться домой к восьми часам.", "вариант 1", "тот же", "")],
    },
}


def sha256_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def dialogue(messages: list[dict]) -> str:
    return "\n".join(f"{m['author']}: {m['text']}" for m in messages)


def row_cells(kind: str, entry: dict) -> tuple[str, ...]:
    """Текстовые ячейки строки после № — то, что человек читает."""
    if kind == "audit":
        return (dialogue(entry["messages"]),)
    return (dialogue(entry["left"]), dialogue(entry["right"]))


# --- render ----------------------------------------------------------------------------------

def render_workbook(kind: str, packet: dict, packet_sha256: str, person: str):
    """openpyxl импортируется здесь, а не на уровне модуля: core tooling его не видит."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    spec = KINDS[kind]
    wb = Workbook()

    # --- Инструкция ---
    ws_i = wb.active
    ws_i.title = "Инструкция"
    ws_i.column_dimensions["A"].width = 110
    ws_i["A1"] = "Инструкция"
    ws_i["A1"].font = Font(name=FONT_NAME, bold=True, size=13)
    row = 3
    for line in spec["instructions"]:
        c = ws_i.cell(row=row, column=1, value=line)
        c.font = Font(name=FONT_NAME, size=11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
    row += 1
    ws_i.cell(row=row, column=1, value="Заполняются только жёлтые ячейки на листе «Оценка». Пример заполнения (не является частью задания):").font = Font(name=FONT_NAME, bold=True, size=11)
    row += 1
    for col, header in enumerate(spec["headers"][2:], start=1):
        c = ws_i.cell(row=row, column=col, value=header)
        c.font = Font(name=FONT_NAME, bold=True, size=10)
        c.fill = PatternFill("solid", fgColor=FILL_HEADER)
    for example in spec["example"]:
        row += 1
        for col, value in enumerate(example, start=1):
            c = ws_i.cell(row=row, column=col, value=value)
            c.font = Font(name=FONT_NAME, size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if col >= len(row_cells(kind, packet[spec["packet_key"]][0])) + 1:
                c.fill = PatternFill("solid", fgColor=FILL_ANSWER)
    for col in range(2, len(spec["headers"]) - 1):
        ws_i.column_dimensions[get_column_letter(col)].width = 26
    ws_i.protection.sheet = True

    # --- Оценка ---
    ws = wb.create_sheet("Оценка")
    headers = spec["headers"]
    for col, header in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=header)
        c.font = Font(name=FONT_NAME, bold=True, size=11)
        c.fill = PatternFill("solid", fgColor=FILL_HEADER)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.column_dimensions["A"].hidden = True
    ws.column_dimensions["B"].width = 5
    n_text = len(row_cells(kind, packet[spec["packet_key"]][0]))
    for col in range(3, 3 + n_text):
        ws.column_dimensions[get_column_letter(col)].width = 60
    for col in spec["choice_columns"]:
        ws.column_dimensions[get_column_letter(col)].width = 24
    for col in spec["text_columns"]:
        ws.column_dimensions[get_column_letter(col)].width = 36
    ws.freeze_panes = "C2"

    validations: dict[int, DataValidation] = {}
    for col, (field, labels) in spec["choice_columns"].items():
        dv = DataValidation(type="list", formula1='"' + ",".join(labels.values()) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle="Только из списка",
                            error="Выберите одно из значений списка.")
        ws.add_data_validation(dv)
        validations[col] = dv
    na_dv = DataValidation(type="list", formula1='"неприменимо"', allow_blank=False, showErrorMessage=True,
                           errorTitle="Неприменимо", error="Для одиночного сообщения этот вопрос не задаётся.")
    ws.add_data_validation(na_dv)

    unlocked = Protection(locked=False)
    for index, entry in enumerate(packet[spec["packet_key"]], start=1):
        r = index + 1
        ws.cell(row=r, column=1, value=entry[spec["row_id_key"]]).font = Font(name=FONT_NAME, size=8)
        ws.cell(row=r, column=2, value=index).font = Font(name=FONT_NAME, size=10)
        for offset, text in enumerate(row_cells(kind, entry)):
            c = ws.cell(row=r, column=3 + offset, value=text)
            c.font = Font(name=FONT_NAME, size=11)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        single = kind == "audit" and len(entry["messages"]) == 1
        for col, (field, labels) in spec["choice_columns"].items():
            c = ws.cell(row=r, column=col)
            c.font = Font(name=FONT_NAME, size=11)
            c.alignment = Alignment(vertical="top")
            if field == "invisible_context" and single:
                c.value = labels["not_applicable"]
                na_dv.add(c)              # выбор невозможен, ячейка остаётся заблокированной
            else:
                c.fill = PatternFill("solid", fgColor=FILL_ANSWER)
                c.protection = unlocked
                validations[col].add(c)
        for col in spec["text_columns"]:
            c = ws.cell(row=r, column=col)
            c.font = Font(name=FONT_NAME, size=11)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.fill = PatternFill("solid", fgColor=FILL_ANSWER)
            c.protection = unlocked
        ws.row_dimensions[r].height = max(30, 15 * (1 + max(t.count("\n") for t in row_cells(kind, entry))))
    ws.protection.sheet = True

    # --- _meta (скрыт) ---
    ws_m = wb.create_sheet("_meta")
    for r, (k, v) in enumerate([
        ("workbook_schema", WORKBOOK_SCHEMA), ("kind", kind), ("person", person),
        ("packet_schema", packet["schema_version"]), ("packet_sha256", packet_sha256),
        ("n_rows", len(packet[spec["packet_key"]])),
    ], start=1):
        ws_m.cell(row=r, column=1, value=k)
        ws_m.cell(row=r, column=2, value=v)
    ws_m.sheet_state = "hidden"
    ws_m.protection.sheet = True
    return wb


def cmd_render(kind: str, base: Path, person: str | None, force: bool) -> int:
    spec = KINDS[kind]
    manifest = json.loads((base / spec["manifest"]).read_text(encoding="utf-8"))
    persons = [person] if person else list(manifest[spec["persons_key"]])
    out_dir = base / "xlsx"
    out_dir.mkdir(exist_ok=True)
    for pid in persons:
        packet_path = base / manifest["packets_dir"] / f"{pid}.json"
        if not packet_path.exists():
            print(f"REFUSED: {packet_path} not found — сначала build пакетов", file=sys.stderr)
            return 1
        target = out_dir / f"{pid}.xlsx"
        if target.exists() and not force:
            print(f"REFUSED: {target} already exists — выданный файл не перерисовывается; "
                  "--force только до выдачи, после — процедура invalidation", file=sys.stderr)
            return 1
        packet = json.loads(packet_path.read_text(encoding="utf-8"))
        wb = render_workbook(kind, packet, sha256_file(packet_path), pid)
        wb.save(target)
        print(f"{pid}: {target.name}, {len(packet[spec['packet_key']])} rows, sha256 {sha256_file(target)[:12]}… "
              f"(запишите полный sha256 файла, который реально отправите)")
    print("OK: workbooks rendered — один файл = один человек, второму не переиспользовать")
    return 0


# --- collect ---------------------------------------------------------------------------------

def collect_workbook(kind: str, wb, packet: dict, packet_sha256: str, person: str) -> tuple[list[dict], list[str]]:
    """Возвращает (responses, issues). Любой issue — отказ; JSONL не пишется."""
    spec = KINDS[kind]
    issues: list[str] = []
    if "_meta" not in wb.sheetnames or "Оценка" not in wb.sheetnames:
        return [], ["structure: листы «Оценка» и «_meta» обязаны существовать под своими именами"]
    meta = {row[0]: row[1] for row in wb["_meta"].iter_rows(min_row=1, max_col=2, values_only=True) if row[0]}
    if meta.get("workbook_schema") != WORKBOOK_SCHEMA:
        issues.append(f"meta: workbook_schema {meta.get('workbook_schema')!r} != {WORKBOOK_SCHEMA}")
    if meta.get("kind") != kind:
        issues.append(f"meta: kind {meta.get('kind')!r} != {kind!r}")
    if meta.get("person") != person:
        issues.append(f"meta: книга выдана {meta.get('person')!r}, а собирается как {person!r} — один файл = один человек")
    if meta.get("packet_sha256") != packet_sha256:
        issues.append("meta: packet_sha256 не совпадает с пакетом на диске — книга сделана из другого пакета (или пакет менялся после выдачи)")
    if issues:
        return [], issues

    ws = wb["Оценка"]
    header = tuple(ws.cell(row=1, column=c).value for c in range(1, len(spec["headers"]) + 1))
    if header != spec["headers"]:
        issues.append(f"structure: заголовки изменены: {header}")
        return [], issues

    expected = {e[spec["row_id_key"]]: e for e in packet[spec["packet_key"]]}
    label_to_code = {col: {label: code for code, label in labels.items()} for col, (_, labels) in spec["choice_columns"].items()}
    responses: list[dict] = []
    seen: dict[str, int] = {}
    for r in range(2, ws.max_row + 1):
        rid = ws.cell(row=r, column=1).value
        if rid is None and all(ws.cell(row=r, column=c).value in (None, "") for c in range(2, len(spec["headers"]) + 1)):
            continue  # пустая строка в хвосте
        if rid not in expected:
            issues.append(f"row {r}: неизвестный или удалённый id {rid!r}")
            continue
        seen[rid] = seen.get(rid, 0) + 1
        for offset, text in enumerate(row_cells(kind, expected[rid])):
            if (ws.cell(row=r, column=3 + offset).value or "") != text:
                issues.append(f"row {r}: текст переписки изменён")
        response = {"schema_version": spec["response_schema"], spec["response_id_key"]: rid, spec["person_key"]: person}
        for col, (field, _) in spec["choice_columns"].items():
            raw = ws.cell(row=r, column=col).value
            label = str(raw).strip().lower() if raw is not None else ""
            code = label_to_code[col].get(label)
            if code is None:
                issues.append(f"row {r}: «{spec['headers'][col - 1]}» = {raw!r} — не из списка")
            response[field] = code
        for col, field in spec["text_columns"].items():
            raw = ws.cell(row=r, column=col).value
            text = str(raw).strip() if raw is not None else ""
            if field == "note":
                if text:
                    response["note"] = text
            else:
                response[field] = text
        responses.append(response)
    for rid, count in seen.items():
        if count > 1:
            issues.append(f"id {rid} встречается {count} раза — строка продублирована")
    for rid in expected:
        if rid not in seen:
            issues.append(f"id {rid} отсутствует — строка удалена")
    return responses, issues


def cmd_collect(kind: str, base: Path, person: str, returned: Path) -> int:
    from openpyxl import load_workbook

    spec = KINDS[kind]
    manifest = json.loads((base / spec["manifest"]).read_text(encoding="utf-8"))
    packet_path = base / manifest["packets_dir"] / f"{person}.json"
    packet = json.loads(packet_path.read_text(encoding="utf-8"))
    out_path = base / manifest["responses_dir"] / f"{person}.jsonl"
    if out_path.exists():
        print(f"REFUSED: {out_path} already exists — собранный слой не перезаписывается никогда", file=sys.stderr)
        return 1
    if not returned.exists():
        print(f"REFUSED: {returned} not found", file=sys.stderr)
        return 1
    returned_sha = sha256_file(returned)
    kept = base / manifest["responses_dir"] / "returned" / f"{person}.xlsx"
    if kept.exists() and sha256_file(kept) != returned_sha:
        print(f"REFUSED: {kept} уже сохранён с другим sha256 — вторая версия ответов не принимается", file=sys.stderr)
        return 1

    wb = load_workbook(returned)
    responses, issues = collect_workbook(kind, wb, packet, sha256_file(packet_path), person)
    if not issues:
        mapping = json.loads((base / manifest["packet_map_dir"] / f"{person}.json").read_text(encoding="utf-8"))["map"]
        if kind == "audit":
            from metrics.naturalness_audit import validate_responses
        else:
            from metrics.naturalness_ab import validate_responses
        issues = validate_responses(responses, mapping, person)
    if issues:
        print("RETURNED WORKBOOK REJECTED:", file=sys.stderr)
        for issue in issues:
            print(f"- {issue}", file=sys.stderr)
        return 1

    kept.parent.mkdir(parents=True, exist_ok=True)
    if not kept.exists():
        shutil.copyfile(returned, kept)
    out_path.write_text("\n".join(json.dumps(r, ensure_ascii=False) for r in responses) + "\n", encoding="utf-8")
    print(f"OK: {len(responses)} responses → {out_path}")
    print(f"returned xlsx kept as {kept} sha256 {returned_sha}")
    print(f"jsonl sha256 {sha256_file(out_path)}  (lock: закоммитьте оба файла как есть)")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    render = sub.add_parser("render", help="packet JSON → per-person .xlsx")
    render.add_argument("--kind", choices=sorted(KINDS), required=True)
    render.add_argument("--dir", type=Path, required=True, help="каталог аудита или A/B")
    render.add_argument("--person", help="один человек; по умолчанию — все из manifest")
    render.add_argument("--force", action="store_true", help="перерисовать существующий файл (ТОЛЬКО до выдачи)")
    collect = sub.add_parser("collect", help="returned .xlsx → response JSONL")
    collect.add_argument("--kind", choices=sorted(KINDS), required=True)
    collect.add_argument("--dir", type=Path, required=True)
    collect.add_argument("--person", required=True)
    collect.add_argument("--returned", type=Path, required=True, help="файл, который прислал человек, как получен")
    args = parser.parse_args()
    if args.command == "render":
        return cmd_render(args.kind, args.dir, args.person, args.force)
    return cmd_collect(args.kind, args.dir, args.person, args.returned)


if __name__ == "__main__":
    sys.exit(main())
